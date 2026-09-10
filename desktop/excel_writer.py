"""
Engine untuk tulis data borang -> cell Excel yang tepat, ikut TEMPLATE_LAYOUT.
Guna openpyxl supaya formula & formatting asal template dikekalkan.
"""
import base64
import io
import re
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

from config import (
    CATEGORIES, FIELDS, TEMPLATE_LAYOUT, TIME_FIELDS, HEADER_CELLS,
    fields_for, template_key_for, category_name, subtype_name,
)
from templates_data import TEMPLATES


def _col_offset(letter, offset):
    idx = column_index_from_string(letter) + offset
    return get_column_letter(idx)


def _resolve_writable_cell(ws, cell_ref):
    """Kalau cell_ref adalah bahagian DALAM satu merge (bukan anchor/top-left),
    pulangkan cell ANCHOR merge tersebut supaya boleh ditulis (openpyxl hanya
    benarkan tulis pada anchor). Kalau bukan sebahagian merge, pulangkan cell asal."""
    cell = ws[cell_ref]
    if cell.__class__.__name__ == "MergedCell":
        for merged_range in ws.merged_cells.ranges:
            if cell.coordinate in merged_range:
                return ws.cell(row=merged_range.min_row, column=merged_range.min_col)
    return cell


def _set_cell(ws, cell_ref, value, field_name=None):
    """Tulis value ke cell_ref, dengan auto-kesan nombor & medan masa.
    Kalau cell_ref sebenarnya sebahagian cell yang di-merge (bukan anchor —
    kuirk template yang kadang berlaku), auto redirect ke anchor supaya
    tidak crash (nota: kalau >1 field kongsi anchor yang sama, nilai
    terakhir ditulis akan menang)."""
    if value is None or value == "":
        return
    cell = _resolve_writable_cell(ws, cell_ref)
    # Medan masa (format "HH:MM") -> objek time Excel sebenar
    if field_name and field_name in TIME_FIELDS:
        m = re.match(r"^(\d{1,2}):(\d{2})$", str(value).strip())
        if m:
            h, mi = int(m.group(1)), int(m.group(2))
            cell.value = datetime(1900, 1, 1, h, mi).time()
            cell.number_format = "hh:mm"
            return
    # Auto-kesan nombor
    s = str(value).strip()
    try:
        if s != "" and not re.match(r"^0\d", s):
            num = float(s)
            cell.value = num
            return
    except ValueError:
        pass
    cell.value = s


def _load_template_ws(tpl_key):
    """Muatkan template (base64) sebagai workbook baharu, pulangkan (wb, ws)."""
    b64 = TEMPLATES.get(tpl_key)
    if not b64:
        return None, None
    data = base64.b64decode(b64)
    wb = load_workbook(io.BytesIO(data))
    ws = wb[wb.sheetnames[0]]
    return wb, ws


def _unique_sheet_name(wb, base_name):
    # Buang aksara yang tak dibenarkan Excel dalam nama sheet
    base_name = re.sub(r'[\\/*\[\]:?]', '-', base_name)
    name = base_name[:31] or "Sheet"
    counter = 1
    while name in wb.sheetnames:
        suffix = f" ({counter})"
        name = (base_name[: 31 - len(suffix)] + suffix)
        counter += 1
    return name


def _write_header(ws, tpl_key, project_name, building_name, location):
    cells = HEADER_CELLS.get(tpl_key)
    if not cells:
        return
    if cells.get("Project"):
        _set_cell(ws, cells["Project"], project_name)
    if cells.get("Building"):
        _set_cell(ws, cells["Building"], building_name)
    if cells.get("Location"):
        _set_cell(ws, cells["Location"], location)


def build_sheet_for_category(cat_id, sub_id, rows, project_name="", building_name="", location=""):
    """
    Bina 1 (atau lebih, utk grid-per-entry) worksheet siap diisi bagi
    kategori:sub tertentu. Pulangkan senarai [(sheet_name_suggestion, ws_or_wb_pair)].
    Untuk mode biasa: pulangkan [(name, ws)] — 1 sheet.
    Untuk grid-per-entry (AHU): pulangkan [(name1, ws1), (name2, ws2), ...] — 1 sheet per entri.
    """
    layout_key = f"{cat_id}:{sub_id}"
    layout = TEMPLATE_LAYOUT.get(layout_key)
    fields = fields_for(cat_id, sub_id)
    tpl_key = template_key_for(cat_id, sub_id)

    cats = [c for c in CATEGORIES if c["id"] == cat_id]
    sub_count = len(cats[0]["subTypes"]) if cats else 1
    short_cat_name = "ACMV" if cat_id == "acmv" else category_name(cat_id)
    raw_name = f"{short_cat_name} - {subtype_name(cat_id, sub_id)}" if sub_count > 1 else short_cat_name

    # ---- grid-per-entry (AHU): 1 sheet PENUH per entri ----
    if layout and layout.get("mode") == "grid-per-entry":
        results = []
        for i, row in enumerate(rows):
            wb, ws = _load_template_ws(tpl_key)
            if ws is None:
                continue
            _write_header(ws, tpl_key, project_name, building_name, location)
            for field_name, cell_ref in layout["cells"].items():
                _set_cell(ws, cell_ref, row.get(field_name), field_name)
            suffix = f" - {row.get('AHU Tag No.')}" if row.get("AHU Tag No.") else (f" #{i+1}" if len(rows) > 1 else "")
            results.append((raw_name + suffix, ws))
        return results

    # ---- mod lain: 1 sheet sahaja ----
    wb, ws = _load_template_ws(tpl_key)
    if ws is not None:
        _write_header(ws, tpl_key, project_name, building_name, location)
    if ws is None:
        # fallback: sheet kosong biasa (tiada template)
        wb = Workbook()
        ws = wb.active
        ws.append(fields)
        for row in rows:
            ws.append([row.get(f, "") for f in fields])
        return [(raw_name, ws)]

    if layout and layout.get("mode") == "table":
        skip_rows = set(layout.get("skip_rows", []))
        r = layout["start_row"]
        for i, row in enumerate(rows):
            while r in skip_rows:
                r += 1
            if layout.get("serial_column"):
                ws[f"{layout['serial_column']}{r}"] = i + 1
            for col_letter, field_name in zip(layout["columns"], fields):
                if not col_letter:
                    continue
                _set_cell(ws, f"{col_letter}{r}", row.get(field_name), field_name)
            r += 1

    elif layout and layout.get("mode") == "linked-tables":
        link_field = layout["link_field"]
        for i, row in enumerate(rows):
            link_val = row.get(link_field, "")
            field_cursor = 1  # fields[0] ialah link_field (global, sudah digunakan)
            for table in layout["tables"]:
                r = table["start_row"] + i
                cols = table["columns"]
                # Lajur pertama jadual = link field (Pump/Lift ID) — auto isi
                _set_cell(ws, f"{cols[0]}{r}", link_val, link_field)
                data_count = table["data_count"]
                remaining_fields = fields[field_cursor: field_cursor + data_count]
                for col_letter, field_name in zip(cols[1:], remaining_fields):
                    _set_cell(ws, f"{col_letter}{r}", row.get(field_name), field_name)
                field_cursor += data_count

    elif layout and layout.get("mode") == "mixed":
        gi = layout["general_info"]
        tb = layout["table"]
        if rows:
            for idx in range(gi["count"]):
                r = gi["start_row"] + idx
                field_name = fields[idx]
                _set_cell(ws, f"{gi['column']}{r}", rows[0].get(field_name), field_name)
        for i, row in enumerate(rows):
            r = tb["start_row"] + i
            for col_letter, field_name in zip(tb["columns"], fields[gi["count"]:]):
                if not col_letter:
                    continue
                _set_cell(ws, f"{col_letter}{r}", row.get(field_name), field_name)

    elif layout and layout.get("mode") == "columns-per-entry":
        for entry_idx, row in enumerate(rows):
            col = _col_offset(layout["start_col"], entry_idx)
            for idx, field_name in enumerate(fields):
                r = layout["start_row"] + idx
                _set_cell(ws, f"{col}{r}", row.get(field_name), field_name)

    else:
        # tiada layout khusus -> append lepas row terakhir template
        start_row = ws.max_row + 1 if ws.max_row > 1 else 2
        for i, row in enumerate(rows):
            r = start_row + i
            for c, field_name in enumerate(fields):
                _set_cell(ws, f"{get_column_letter(c+1)}{r}", row.get(field_name), field_name)

    return [(raw_name, ws)]


def export_floor_to_workbook(wb, floor_data, building_name="", floor_name="", project_name="", location=""):
    """
    Isi 1 Workbook (openpyxl) dengan semua sheet kategori bagi SATU floor.
    floor_data: dict {"catId:subId": [row_dict, ...]}
    """
    any_data = False
    for cat in CATEGORIES:
        for sub in cat["subTypes"]:
            key = f"{cat['id']}:{sub['id']}"
            rows = floor_data.get(key, [])
            if not rows:
                continue
            sheets = build_sheet_for_category(cat["id"], sub["id"], rows, project_name, building_name, location)
            for name, ws in sheets:
                any_data = True
                sheet_name = _unique_sheet_name(wb, name)
                # Salin worksheet (dari workbook sumbernya) ke workbook destinasi
                _copy_ws_into(wb, ws, sheet_name)
    return any_data


def _copy_ws_into(dest_wb, src_ws, sheet_name):
    """Salin src_ws (dari workbook lain) sepenuhnya ke dest_wb sebagai sheet baharu,
    termasuk nilai, formula, format nombor, merged cells, dan lebar lajur."""
    new_ws = dest_wb.create_sheet(title=sheet_name)
    for row in src_ws.iter_rows():
        for cell in row:
            new_cell = new_ws.cell(row=cell.row, column=cell.column, value=cell.value)
            if cell.has_style:
                new_cell.font = cell.font.copy()
                new_cell.border = cell.border.copy()
                new_cell.fill = cell.fill.copy()
                new_cell.number_format = cell.number_format
                new_cell.protection = cell.protection.copy()
                new_cell.alignment = cell.alignment.copy()
    for merged_range in src_ws.merged_cells.ranges:
        new_ws.merge_cells(str(merged_range))
    for col_letter, dim in src_ws.column_dimensions.items():
        new_ws.column_dimensions[col_letter].width = dim.width
    for row_idx, dim in src_ws.row_dimensions.items():
        new_ws.row_dimensions[row_idx].height = dim.height
    return new_ws


def export_project_to_one_workbook(buildings, project_name=""):
    """
    Compile SEMUA bangunan + floor ke DALAM SATU fail xlsx.
    buildings: list of {"name":..., "location":..., "floors":[{"name":..., "data":{...}}]}
    Sheet dinamakan: "<Floor> - <Kategori>" (dipendekkan supaya <=31 aksara),
    dan kalau berbilang bangunan, prefix nama bangunan turut disertakan.
    """
    wb = Workbook()
    wb.remove(wb.active)  # buang default "Sheet"
    any_data = False

    multi_building = len(buildings) > 1
    for building in buildings:
        b_name = building.get("name", "Building")
        b_location = building.get("location", "")
        b_project = building.get("project") or project_name
        for floor in building.get("floors", []):
            f_name = floor.get("name", "Floor")
            floor_data = floor.get("data", {})
            for cat in CATEGORIES:
                for sub in cat["subTypes"]:
                    key = f"{cat['id']}:{sub['id']}"
                    rows = floor_data.get(key, [])
                    if not rows:
                        continue
                    sheets = build_sheet_for_category(cat["id"], sub["id"], rows, b_project, b_name, b_location)
                    for name, ws in sheets:
                        any_data = True
                        prefix = f"{b_name} " if multi_building else ""
                        full_name = f"{prefix}{f_name} - {name}"
                        sheet_name = _unique_sheet_name(wb, full_name)
                        _copy_ws_into(wb, ws, sheet_name)

    if not any_data:
        wb.create_sheet("(Tiada Data)")
    return wb
